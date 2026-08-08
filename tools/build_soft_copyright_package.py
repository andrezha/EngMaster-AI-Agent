from __future__ import annotations

import datetime as _dt
import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "soft_copyright_engmaster_v1_package"
DOCS = OUT / "01_正式文档"
SOURCE = OUT / "02_源代码材料"
CHECK = OUT / "03_待填写与截图清单"
ARCHIVE = OUT / "99_草稿备份"
SCREENSHOT_SRC = ROOT / "soft_copyright_engmaster_materials" / "screenshots"
SCREENSHOT_OUT = DOCS / "截图"

TODAY = "2026-08-02"
SOFTWARE_NAME = "EngMaster英语词汇分级学习平台软件"
SHORT_NAME = "EngMaster英语词汇分级学习平台"
VERSION = "V1.0"
VOCABULARY_SCOPE_NOTICE = (
    "本软件词汇及学习内容范围主要参考国家英语课程标准、相关英语考试大纲及公开考试要求，"
    "由开发者结合不同学习阶段的实际需求整理编排。本软件为个人开发的英语学习辅助工具，"
    "并非教育主管部门、学校或考试机构官方指定软件。"
)


def _set_cell_text(cell, text: str):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(10.5)


def _set_doc_defaults(doc: Document):
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)

    styles = doc.styles
    styles["Normal"].font.name = "宋体"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    styles["Normal"].font.size = Pt(10.5)
    for name, size in [("Title", 18), ("Heading 1", 15), ("Heading 2", 13), ("Heading 3", 11)]:
        style = styles[name]
        style.font.name = "宋体"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        style.font.size = Pt(size)


def _add_page_number(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def _add_footer(doc: Document):
    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = 1
    footer.add_run(f"{SOFTWARE_NAME} {VERSION}    第 ")
    _add_page_number(footer)
    footer.add_run(" 页")


def _p(doc: Document, text: str = "", style: str | None = None):
    paragraph = doc.add_paragraph(style=style)
    if text:
        run = paragraph.add_run(text)
        run.font.name = "宋体"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    return paragraph


def _bullets(doc: Document, items: list[str]):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(item)
        run.font.name = "宋体"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        run.font.size = Pt(10.5)


def _numbered(doc: Document, items: list[str]):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        run = p.add_run(item)
        run.font.name = "宋体"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        run.font.size = Pt(10.5)


def _screenshot(doc: Document, filename: str, caption: str):
    path = SCREENSHOT_SRC / filename
    if not path.exists():
        _p(doc, f"【待补截图：{caption}】")
        return
    paragraph = doc.add_paragraph()
    paragraph.alignment = 1
    paragraph.add_run().add_picture(str(path), width=Cm(15.5))
    caption_p = doc.add_paragraph(f"图：{caption}")
    caption_p.alignment = 1


def _legacy_build_user_manual():
    doc = Document()
    _set_doc_defaults(doc)
    _add_footer(doc)

    doc.add_heading(f"{SOFTWARE_NAME} {VERSION}", 0)
    doc.add_heading("用户操作说明书", 1)
    _p(doc, f"文档日期：{TODAY}")
    _p(doc, "说明：本文档为软件著作权登记操作说明书，界面截图取自当前 V1.0 软件实际运行版本。")
    doc.add_page_break()

    doc.add_heading("一、软件简介", 1)
    _p(
        doc,
        f"{SOFTWARE_NAME}是一款运行于 Windows 桌面环境的英语学习辅助软件。软件面向高中英语学习和高考英语备考场景，"
        "提供高考 3800 词闯关、词汇表查看、自主词库、短语与不规则动词练习、专项题型模拟练习、整卷模拟练习、"
        "参考解析、个人错词及错题打印表生成、用户须知展示和授权激活等功能。",
    )
    _p(
        doc,
        "软件中的练习内容用于模拟训练和学习复盘，不作为官方考试题目、考试预测或学习结果承诺。"
        "用户应结合教材、课堂内容、教师指导和正式考试要求进行学习。",
    )
    _p(doc, VOCABULARY_SCOPE_NOTICE)

    doc.add_heading("二、运行环境", 1)
    table = doc.add_table(rows=5, cols=2)
    table.style = "Table Grid"
    rows = [
        ("软件名称", SOFTWARE_NAME),
        ("软件简称", SHORT_NAME),
        ("版本号", VERSION),
        ("推荐运行环境", "Windows 10 / Windows 11 64 位系统"),
        ("运行方式", "纯单机运行；授权激活采用单机激活码和本地授权校验。"),
    ]
    for row, (k, v) in zip(table.rows, rows):
        _set_cell_text(row.cells[0], k)
        _set_cell_text(row.cells[1], v)

    doc.add_heading("三、软件启动与授权", 1)
    _numbered(
        doc,
        [
            "打开软件所在目录，双击可执行程序启动软件。",
            "首次启动或授权文件不存在时，软件会显示用户须知与激活窗口。",
            "用户阅读并确认用户须知后，在激活窗口输入单机激活码。",
            "软件根据单机激活码和本机机器码进行本地授权校验。",
            "授权成功后，软件保存本地授权文件，并进入主界面。",
            "单机激活码原则上仅限绑定一台电脑。更换电脑、重装系统、更换主板或系统环境变化导致机器码改变时，可能需要重新授权或联系购买渠道处理。",
        ],
    )
    _screenshot(doc, "08_用户须知与免责声明.png", "用户须知与免责声明界面")
    _screenshot(doc, "09_单机激活界面.png", "单机激活界面")
    _screenshot(doc, "01_主界面_词汇闯关.png", "软件主界面")

    doc.add_heading("四、高考 3800 词闯关", 1)
    _numbered(
        doc,
        [
            "在左侧导航栏选择高考 3800 词闯关功能。",
            "选择常规词汇闯关、错词闯关或自主录入词汇闯关。",
            "查看页面显示的单词或释义提示，在输入框中填写答案。",
            "点击确认按钮或按回车提交答案。",
            "系统显示答题结果，并根据答题情况记录错题。",
        ],
    )
    _screenshot(doc, "01_主界面_词汇闯关.png", "高考 3800 词闯关界面")

    doc.add_heading("五、词汇表与资料查看", 1)
    _numbered(
        doc,
        [
            "在左侧导航栏选择高考 3800 词汇表功能。",
            "在列表中查看内置词汇、错题词汇、短语或不规则动词。",
            "使用搜索框查找指定内容。",
            "使用上一页、下一页进行分页查看。",
            "根据复习需要选择隐藏英文或隐藏中文。",
        ],
    )
    _screenshot(doc, "02_高考3800词汇表.png", "高考 3800 词汇表界面")

    doc.add_heading("六、自主词库管理", 1)
    _numbered(
        doc,
        [
            "进入自主词库页面。",
            "在英文单词输入框中输入单词，在中文释义输入框中输入解释。",
            "点击添加按钮保存词汇。",
            "在词汇表中查看、编辑或删除已添加内容。",
            "可选择中英对照、只看英语或只看中文进行背诵。",
            "自主词库可用于个人化词汇复习。",
        ],
    )
    _screenshot(doc, "03_自主词库.png", "自主词库管理界面")

    doc.add_heading("七、短语与不规则动词练习", 1)
    _numbered(
        doc,
        [
            "进入短语与不规则动词挑战功能。",
            "选择短语练习或不规则动词练习。",
            "短语列表可切换核心、扩展及中英文显示方式；不规则动词列表可隐藏变化形式辅助背诵。",
            "根据页面提示输入答案并提交。",
            "系统显示答题结果，并记录错误项目。",
            "用户可进入相应列表查看和复习错题。",
        ],
    )
    _screenshot(doc, "04_短语与不规则动词练习.png", "短语与不规则动词练习界面")
    _screenshot(doc, "05_短语与不规则动词列表.png", "短语与不规则动词列表界面")

    doc.add_heading("八、专项题型模拟练习", 1)
    _numbered(
        doc,
        [
            "进入高考题型模拟练习功能。",
            "选择阅读理解、完形填空、七选五或语法填空题型。",
            "系统载入本地模拟练习资源并生成答题界面。",
            "用户完成选择或填空后提交答案。",
            "系统显示得分、正确答案和参考解析。",
        ],
    )
    _screenshot(doc, "06_高考题型模拟练习.png", "专项题型模拟练习界面")

    doc.add_heading("九、整卷模拟练习", 1)
    _numbered(
        doc,
        [
            "进入整卷模拟练习功能。",
            "系统从本地模拟试卷资源中载入一份练习文本。",
            "用户按题型顺序完成作答。",
            "提交后系统展示得分、答案和参考解析。",
            "用户可根据结果进行复盘。",
        ],
    )
    _screenshot(doc, "07_整卷模拟练习.png", "整卷模拟练习界面")

    doc.add_heading("十、错词打印表生成", 1)
    _numbered(
        doc,
        [
            "进入词汇表并打开错词表。",
            "选择个人错词、短语错题或不规则动词错题。",
            "选择打印表或默写打印表复习模式。",
            "点击生成错词打印表按钮。",
            "系统生成对应 PDF 文档，用户可打开查看或打印。",
        ],
    )
    _p(doc, "说明：打印功能仅输出个人错词、短语错题和不规则动词错题，完整词表及完整学习资料暂不提供打印。")
    _p(doc, "截图位置：错词导出入口、保存窗口、导出后的 PDF 页眉、水印和页脚。")

    doc.add_heading("十一、本地数据说明", 1)
    _p(
        doc,
        "软件的错词记录、自主词库、学习记录和使用配置等数据主要保存在用户本机。用户删除软件、清理系统文件、"
        "重装系统、更换设备、磁盘损坏或误删文件，可能导致本地学习数据丢失。请用户自行妥善备份重要学习数据。"
        "因上述原因造成的本地学习数据丢失，不属于软件质量问题。",
    )

    doc.add_heading("十二、常见问题", 1)
    _p(doc, "软件内可通过左侧菜单的“常见问题”按钮直接打开本页内容。")
    qa = [
        ("词表只显示英语或只显示中文", "请检查是否选择了“只看英语”或“只看中文”；需要恢复时点击“中英对照”。"),
        ("软件无法启动", "请确认运行环境是否为 Windows 10 / Windows 11 64 位系统，并确认软件文件未被删除或移动。"),
        ("单机激活码无法通过", "请确认单机激活码是否输入完整，并联系购买渠道确认授权状态。"),
        ("更换电脑后单机激活码无法使用", "单机激活码原则上仅限绑定一台电脑。如更换电脑、重装系统、更换主板或机器码变化，请联系购买渠道处理。"),
        ("学习数据丢失", "请检查本地用户数据目录是否被清理或覆盖。重要学习数据建议由用户自行备份。"),
    ]
    for q, a in qa:
        doc.add_heading(q, 2)
        _p(doc, a)

    doc.add_heading("十三、主要界面说明", 1)
    _bullets(
        doc,
        [
            "用户须知与授权激活截图",
            "主界面截图",
            "高考 3800 词闯关页面截图",
            "高考 3800 词汇表页面截图",
            "自主词库页面截图",
            "短语与不规则动词页面截图",
            "专项题型模拟练习页面截图",
            "整卷模拟练习页面截图",
            "以上截图均取自当前 V1.0 软件实际运行界面，用于说明各功能模块的操作入口和页面布局。",
        ],
    )

    doc.save(DOCS / f"{SOFTWARE_NAME}_{VERSION}_用户操作说明书.docx")


def build_user_manual():
    doc = Document()
    _set_doc_defaults(doc)
    _add_footer(doc)

    doc.add_heading(f"{SOFTWARE_NAME} {VERSION}", 0)
    doc.add_heading("用户操作说明书", 1)
    _p(doc, f"文档日期：{TODAY}")
    _p(doc, "说明：本文档描述当前 V1.0 实际开放功能，正式截图须来自同一最终发布程序。")
    doc.add_page_break()

    doc.add_heading("一、软件简介", 1)
    _p(
        doc,
        f"{SOFTWARE_NAME}是一款 Windows 桌面端分级英语词汇学习软件。软件通过词表学习、"
        "普通闯关、个人错词表、错词闯关和下一轮复测形成闭环学习流程，并提供自主词库、"
        "短语与不规则动词训练、轮次学习记录以及资料导出功能。",
    )
    _p(
        doc,
        "V1.0 提供初中、高考、大学英语四级、大学英语六级和考研五个级别的30词免费体验。"
        "当前仅高考英语正式版完成并开放购买；其他正式版处于开发阶段，不作为本版本已开放功能。",
    )
    _p(doc, VOCABULARY_SCOPE_NOTICE)

    doc.add_heading("二、运行环境", 1)
    table = doc.add_table(rows=5, cols=2)
    table.style = "Table Grid"
    rows = [
        ("软件名称", SOFTWARE_NAME),
        ("软件简称", SHORT_NAME),
        ("版本号", VERSION),
        ("推荐环境", "Windows 10 / Windows 11 64 位系统"),
        ("运行方式", "单机免安装运行；体验无需激活，高考正式版采用本地单机激活。"),
    ]
    for row, (key, value) in zip(table.rows, rows):
        _set_cell_text(row.cells[0], key)
        _set_cell_text(row.cells[1], value)
    _p(
        doc,
        "当前发布文件未附加商业代码签名。交付 ZIP 中包含独立的《运行前必读》安全指南和"
        "SHA-256 校验值。用户应先确认官方来源并核对校验值，再按指南处理 Windows SmartScreen"
        "或杀毒软件提醒；不应关闭全部实时保护、防火墙或添加宽泛目录排除。",
    )

    doc.add_heading("三、启动与免费体验", 1)
    _numbered(doc, [
        "双击可执行程序启动软件。",
        "未检测到有效授权时，软件直接进入免费体验模式，不要求注册账号。",
        "在体验中心选择初中、高考、四级、六级或考研级别。",
        "每个级别提供独立的30词体验，进度、错词和轮次记录分别保存。",
        "通过快速了解、学习方法和操作指南查看推荐学习流程。",
    ])
    _screenshot(doc, "01_免费体验中心.png", "免费体验中心")
    _screenshot(doc, "02_体验级别选择.png", "五级体验选择区域")
    _screenshot(doc, "03_学习方法指南.png", "学习闭环与操作指南")

    doc.add_heading("四、正式版购买、订单核对与激活", 1)
    _numbered(doc, [
        "在免费体验页面点击“前往购买正式版”，进入英语正式版管理与购买页面。",
        "页面展示各正式版的开放状态；V1.0 仅高考英语正式版开放购买，其他正式版显示“开发中”。",
        "点击统一的“前往淘宝购买”，在淘宝商品页面选择实际开放的商品规格并完成付款。软件不在本地预选购买版本，也不自动读取淘宝订单或付款信息。",
        "普通买家填写淘宝订单号；朋友或测试人员可填写客服提供的登记号。软件不会自动联网验证订单。",
        "软件将订单号和本机识别码组合成客服核验信息，再发送给对应客服。",
        "客服按订单号核对对应登记记录后，再为该机器码提供相应的累计单机激活码。",
        "返回本页阅读并接受用户须知与免责声明，粘贴激活码并完成本地校验。成功后相应正式版显示绿色“已解锁”。",
        "激活码与电脑绑定；更换设备或机器识别信息变化时需联系购买渠道处理。",
    ])
    _screenshot(doc, "08_正式版管理与淘宝购买.png", "统一淘宝购买入口及正式版开放状态")
    _screenshot(doc, "09_用户须知.png", "用户须知与免责声明")
    _screenshot(doc, "10_客服信息与单机激活.png", "机器码复制、订单核对说明及单机激活区域")

    doc.add_heading("五、词汇表学习", 1)
    _numbered(doc, [
        "进入当前级别词汇表，选择内置词汇、个人错词或自主录入词汇。",
        "使用搜索框查找指定单词。",
        "选择中英对照、隐藏英文或隐藏中文等显示方式。",
        "按页浏览和背诵；错词积累后可生成错词纸质学习资料。",
    ])
    _screenshot(doc, "04_高考词汇表.png", "高考英语词汇表")

    doc.add_heading("六、词汇闯关与错词闭环", 1)
    _numbered(doc, [
        "进入普通词汇闯关，根据英文或中文提示输入答案。",
        "答错或掌握不稳定的单词自动加入个人错词表。",
        "返回错词表集中背诵，再进入错词闯关反复验证。",
        "错词连续答对达到规定次数后从当前错词表移出。",
        "重新进行下一轮普通闯关，直到完整一轮不再产生新错词。",
        "在学习记录中查看轮次、用时、答题数量和错词变化。",
    ])
    _screenshot(doc, "05_高考词汇闯关.png", "高考英语词汇闯关")
    _screenshot(doc, "06_错词与学习记录.png", "个人错词与轮次学习记录")

    doc.add_heading("七、自主词库", 1)
    _numbered(doc, [
        "进入自主登记单词页面，输入英文单词和中文释义后保存。",
        "查看、编辑或删除已经登记的个人词汇。",
        "根据背诵需要选择中英对照、只看英语或只看中文。",
        "自主词汇可进入独立闯关，不与内置正式词库混合。",
    ])
    _screenshot(doc, "07_自主登记单词.png", "自主词库管理界面")

    doc.add_heading("八、短语与不规则动词", 1)
    _numbered(doc, [
        "进入短语与不规则动词表，短语可切换核心、扩展及中英文显示方式。",
        "不规则动词可选择完整对照、看原形默变化或只看中文。",
        "进入对应闯关，根据提示填写短语或动词变化形式。",
        "错误项目进入独立错题记录，可集中复习和再次验证。",
        "体验版提供示例内容，高考正式版提供完整学习内容。",
    ])
    _screenshot(doc, "12_短语与不规则动词.png", "短语与不规则动词学习界面")

    doc.add_heading("九、错题资料导出", 1)
    _numbered(doc, [
        "在词汇表中打开个人错词、短语错题或不规则动词错题。",
        "仅针对当前错词或错题选择普通对照、看英默中或看中默英等输出形式。",
        "点击生成错词打印表，系统生成 PDF 并保存到用户指定位置。",
        "打开生成文件进行查看或打印。",
    ])
    _p(doc, "完整词表、完整短语表、不规则动词完整资料及自主录入完整词表暂不提供打印。")
    _screenshot(doc, "13_学习资料导出.png", "错词资料导出界面")

    doc.add_heading("十、本地数据与版本隔离", 1)
    _p(
        doc,
        "错词、自主词库、学习记录和设置主要保存在用户本机。五类体验数据分别保存，"
        "体验数据与高考正式版数据相互隔离。清理系统文件、重装系统、磁盘损坏或误删"
        "可能造成学习数据丢失，用户应自行备份重要记录。",
    )

    doc.add_heading("十一、常见问题", 1)
    _p(doc, "软件内可通过左侧菜单的“常见问题”按钮直接打开本页内容。")
    for question, answer in [
        ("词表只显示英语或只显示中文", "请检查是否选择了“只看英语”或“只看中文”；需要恢复时点击“中英对照”。"),
        ("软件无法启动", "请确认使用 Windows 10 或 Windows 11 64 位系统，并确认程序文件完整。"),
        ("其他正式版为什么不能购买", "V1.0 仅开放高考英语正式版；其他正式版开发中，可使用对应30词体验。"),
        ("客服如何确认购买版本", "普通买家填写淘宝订单号；朋友或测试人员填写客服提供的登记号。客服按订单号核对对应登记记录后处理。"),
        ("激活码无法通过", "请确认订单已付款、订单规格与授权版本一致、激活码输入完整，并联系购买渠道核对。"),
        ("更换电脑后无法使用", "单机激活与电脑绑定，更换设备或机器识别信息变化时请联系购买渠道。"),
        ("学习数据丢失", "检查本机学习数据是否被清理或覆盖，并从用户自行保存的备份恢复。"),
    ]:
        doc.add_heading(question, 2)
        _p(doc, answer)

    doc.add_heading("十二、截图一致性要求", 1)
    _bullets(doc, [
        "截图必须来自同一最终 V1.0 可执行程序。",
        "截图应覆盖免费体验、统一淘宝购买入口、订单核对说明、单机激活、版本状态及主要学习功能。",
        "截图不得包含真实机器码、激活码、订单、个人信息或内部客服工具。",
        "界面中高考正式版应显示可购买或已解锁，其他正式版应显示开发中。",
        "购买流程截图不得出现软件内的“意向版本”或各版本独立购买按钮；版本应由用户在淘宝商品规格中选择。",
    ])

    doc.save(DOCS / f"{SOFTWARE_NAME}_{VERSION}_用户操作说明书.docx")


def build_application_reference():
    doc = Document()
    _set_doc_defaults(doc)
    _add_footer(doc)
    doc.add_heading(f"{SOFTWARE_NAME} {VERSION}", 0)
    doc.add_heading("软著申请信息确认表", 1)
    _p(doc, "说明：本表用于正式填报前核对信息。标注【待填写】的内容需由申请人本人确认。")

    table = doc.add_table(rows=14, cols=2)
    table.style = "Table Grid"
    rows = [
        ("软件全称", SOFTWARE_NAME),
        ("软件简称", SHORT_NAME),
        ("版本号", VERSION),
        ("软件分类", "教育学习类桌面应用软件"),
        ("开发方式", "独立开发"),
        ("权利取得方式", "原始取得"),
        ("著作权人", "【待填写：申请人姓名或主体名称】"),
        ("开发完成日期", "【待填写：请按实际完成日期填写】"),
        ("首次发表状态", "【待填写：未发表 / 已发表】"),
        ("首次发表日期", "【待填写：如未发表则按平台要求填写】"),
        ("运行环境", "Windows 10 / Windows 11 64 位系统"),
        ("编程语言", "Python"),
        ("主要技术", "PySide6 / Qt 图形界面、分级配置、本地文件存储、PDF/文档生成、离线授权校验"),
        ("源程序量参考", "以最终冻结源码重新统计；提交材料只包含自有业务源码。"),
    ]
    for row, (k, v) in zip(table.rows, rows):
        _set_cell_text(row.cells[0], k)
        _set_cell_text(row.cells[1], v)

    doc.add_heading("软件主要功能简述", 1)
    _p(
        doc,
        "本软件提供初中、高考、大学英语四级、大学英语六级和考研五个级别的英语词汇体验，"
        "通过词表学习、普通闯关、个人错词表、错词闯关和轮次复测形成学习闭环，并提供自主词库、"
        "短语与不规则动词训练、学习记录及资料导出。V1.0 仅开放高考英语正式版，采用单机离线授权；"
        "其他正式版仍在开发中，不列为本版本已完成的正式功能。",
    )
    _p(
        doc,
        "正式版购买在外部淘宝商品页面完成，用户在淘宝商品规格中选择版本。软件提供统一购买入口、"
        "订单号填写、客服核验信息复制和本地激活功能，但不自动连接淘宝，也不自动读取或验证订单、付款状态；"
        "销售方依据客户发送的订单号和登记记录人工核对授权范围。",
    )

    doc.add_heading("提交前需申请人确认", 1)
    _bullets(
        doc,
        [
            "软件名称、简称和版本号是否最终采用本表内容。",
            "开发完成日期和首次发表日期是否真实、准确。",
            "申请主体信息是否与证件或营业执照一致。",
            "说明书截图是否来自最终发布版本软件。",
            "源代码材料中是否已排除私钥、账号、测试码、真实用户数据和无关第三方源码。",
            "正式申请表中的功能描述是否只包含当前版本真实可用的功能。",
        ],
    )
    doc.save(DOCS / f"{SOFTWARE_NAME}_{VERSION}_软著申请信息确认表.docx")


def _eligible_source_files() -> list[Path]:
    preferred = [
        "main.py",
        "edition_config.py",
        "experience_guide.py",
        "guide_pages.py",
        "purchase_config.py",
        "purchase_activation_panel.py",
        "version_management.py",
        "trial_center.py",
        "challenge_rounds.py",
        "challenge_history_dialog.py",
        "vocab_module.py",
        "word_list_view.py",
        "phrase_irregular_module.py",
        "self_register_vocab_module.py",
        "pdf_document_generator.py",
        "word_document_generator.py",
        "export_dialog.py",
        "make_word_table.py",
        "text_table_generator.py",
        "utils.py",
        "watermark_modes.py",
    ]
    return [ROOT / p for p in preferred if (ROOT / p).exists()]


def _collect_code_lines() -> list[str]:
    lines: list[str] = []
    for path in _eligible_source_files():
        rel = path.relative_to(ROOT).as_posix()
        lines.append(f"# ===== 文件：{rel} =====")
        text = path.read_text(encoding="utf-8", errors="replace").splitlines()
        for idx, line in enumerate(text, 1):
            lines.append(f"{idx:04d}: {line}")
        lines.append("")
    return lines


def build_source_code_doc():
    lines = _collect_code_lines()
    lines_per_page = 50
    page_count = 60
    required = lines_per_page * page_count
    if len(lines) >= required:
        half = required // 2
        selected = lines[:half] + lines[-half:]
    else:
        selected = lines
        page_count = (len(selected) + lines_per_page - 1) // lines_per_page

    doc = Document()
    _set_doc_defaults(doc)
    _add_footer(doc)
    doc.add_heading(f"{SOFTWARE_NAME} {VERSION}", 0)
    doc.add_heading("源代码提交材料", 1)
    _p(doc, "说明：本文档为软件著作权登记源程序鉴别材料，按源程序前、后各连续 30 页整理，每页 50 行。")
    _p(doc, "整理范围：应用程序自有源码。已排除 venv、build、dist、__pycache__、测试脚本、工具脚本、授权私密材料和第三方库源码。")
    _p(doc, f"生成日期：{TODAY}")
    doc.add_page_break()

    for page in range(page_count):
        doc.add_heading(f"第 {page + 1} 页", 2)
        chunk = selected[page * lines_per_page : (page + 1) * lines_per_page]
        for line in chunk:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            run = p.add_run(line[:150])
            run.font.name = "Consolas"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
            run.font.size = Pt(8)
        if page != page_count - 1:
            doc.add_page_break()

    doc.save(SOURCE / f"{SOFTWARE_NAME}_{VERSION}_源代码提交材料_正式版.docx")


def build_checklists():
    checklist = f"""# 正式提交前待填写与确认清单

生成日期：{TODAY}

## 申请人必须填写

- 著作权人姓名或主体名称：【待填写】
- 身份证号或统一社会信用代码：【待填写】
- 联系地址、电话、邮箱：【待填写】
- 开发完成日期：【待填写】
- 首次发表状态：【待填写：未发表 / 已发表】
- 首次发表日期：【待填写】

## 材料需要人工确认

- 软件全称是否确定为：{SOFTWARE_NAME}
- 软件简称是否确定为：{SHORT_NAME}
- 版本号是否确定为：{VERSION}
- 说明书截图是否来自最终发布版本。
- 截图中是否没有测试码、个人隐私、后台管理信息或乱码。
- 源代码材料是否已排除 RSA 私钥、后台管理信息、真实用户数据和第三方库源码。
- 截图和文档是否明确：V1.0 只开放高考英语正式版。
- 初中、四级、六级和考研是否统一标为“正式版开发中”，没有写成可购买功能。
- 免费体验、统一淘宝购买、订单核对说明、激活、版本管理和学习闭环是否与最终 EXE 一致。
- 文档是否明确：购买版本由用户在淘宝商品规格中选择；软件只组合用户填写的订单号和机器码，不自动连接淘宝或验证付款。
- 文档是否明确：客服以订单号、付款状态和商品规格为准，机器码本身不代表购买版本。
- 页面和截图中是否已经删除软件内“意向版本”及各版本独立购买按钮。
- 用户须知是否包含单机激活绑定、换机处理、本地数据备份和免责说明。

## 最终版需要补拍的界面截图

- `01_免费体验中心.png`
- `02_体验级别选择.png`
- `03_学习方法指南.png`
- `04_高考词汇表.png`
- `05_高考词汇闯关.png`
- `06_错词与学习记录.png`
- `07_自主登记单词.png`
- `08_正式版管理与淘宝购买.png`：只能显示统一淘宝购买入口，不得出现软件内“意向版本”或各版本独立购买按钮。
- `09_用户须知.png`
- `10_客服信息与单机激活.png`：订单号输入框保持为空，机器码必须使用演示值或遮挡，不得包含真实订单、真实激活码或客户信息。
- `12_短语与不规则动词.png`
- `13_学习资料导出.png`

以上截图必须来自同一最终 V1.0 EXE，窗口尺寸和界面字体尽量保持一致。

## 建议最终资料格式

- 软件操作说明书：Word 或 PDF。
- 源代码提交材料：Word 或 PDF，按平台要求上传。
- 申请表：在平台页面填写。
- 身份证明或主体证明：按申请人类型准备。

## 本次已生成文件

- `01_正式文档/{SOFTWARE_NAME}_{VERSION}_用户操作说明书.docx`
- `01_正式文档/{SOFTWARE_NAME}_{VERSION}_软著申请信息确认表.docx`
- `02_源代码材料/{SOFTWARE_NAME}_{VERSION}_源代码提交材料_正式版.docx`
- `03_待填写与截图清单/正式提交前待填写与确认清单.md`
"""
    (CHECK / "正式提交前待填写与确认清单.md").write_text(checklist, encoding="utf-8")

    readme = f"""# {SOFTWARE_NAME} {VERSION} 正式资料包

本目录为软著申请提交前整理包。资料由项目文件整理生成，正式提交前需要申请人本人核对和确认。

## 目录

- `01_正式文档`：操作说明书、申请信息确认表。
- `02_源代码材料`：源代码提交材料正式版。
- `03_待填写与截图清单`：需要申请人补充和确认的事项。
- `99_草稿备份`：Markdown 草稿、质检报告等备份资料。

## 注意

当前生成包先完成文档结构与源代码材料。界面截图必须在程序整体验收和最终 EXE 冻结后重新拍摄，
不得继续使用旧“高中-高考英语单词助手”的截图。正式提交前还需填写身份、日期和发表状态。
"""
    (OUT / "README_先看我.md").write_text(readme, encoding="utf-8")


def backup_drafts():
    draft = ROOT / "soft_copyright_engmaster_materials"
    if draft.exists():
        shutil.copytree(draft, ARCHIVE / "soft_copyright_materials", dirs_exist_ok=True)
    report = ROOT / "出厂质检报告_delivery_report.md"
    if report.exists():
        shutil.copy2(report, ARCHIVE / report.name)


def main():
    if OUT.exists():
        shutil.rmtree(OUT)
    for directory in (DOCS, SOURCE, CHECK, ARCHIVE):
        directory.mkdir(parents=True, exist_ok=True)

    if SCREENSHOT_SRC.exists():
        shutil.copytree(SCREENSHOT_SRC, SCREENSHOT_OUT, dirs_exist_ok=True)

    build_user_manual()
    build_application_reference()
    build_source_code_doc()
    build_checklists()
    backup_drafts()
    print(OUT)


if __name__ == "__main__":
    main()
