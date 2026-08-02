# EngMaster Word document generator
import os
import re  # 必须导入，用于处理“或”、括号及音标符号切分
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL, WD_ROW_HEIGHT_RULE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from utils import _normalize_full_width_to_half_width

# --- 字体配置 ---
FONT_NAME_CN = "微软雅黑"
FONT_NAME_EN = "Times New Roman"  # Roma 字体，确保音标和英文显示正统
FONT_SIZE_INDEX = 8
FONT_SIZE_CONTENT = 10

def set_run_font(run, size, bold=False):
    """
    强制锁定 Roma 字体并兼容中文字体，解决 Windows 下 Word 渲染异常问题
    """
    run.font.size = Pt(size)
    run.bold = bold
    run.font.name = FONT_NAME_EN
    rPr = run._element.get_or_add_rPr()
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), FONT_NAME_EN)
    rFonts.set(qn('w:hAnsi'), FONT_NAME_EN)
    rFonts.set(qn('w:eastAsia'), FONT_NAME_CN)
    rFonts.set(qn('w:hint'), 'default')

    existing = rPr.find(qn('w:rFonts'))
    if existing is not None:
        rPr.remove(existing)
    rPr.append(rFonts)

def generate_word_table(data, output_file, mode="normal", data_type="words"): # Added data_type parameter
    document = Document()

    # 1. 窄边距设置
    section = document.sections[0]
    section.top_margin = Cm(1.2)
    section.bottom_margin = Cm(1.2)
    section.left_margin = Cm(1.0)
    section.right_margin = Cm(1.0)

    # 2. 模式与键名识别
    # en_dictate_cn: 看英语写中文
    # cn_dictate_en: 看中文/发音写英语 (本次优化重点)
    # 初始化默认值，适用于 "normal" 模式和 "words" 数据类型
    v_key, b_key, is_dict = "word", "content", False
    visible_header_text = "英语单词"
    blank_header_text = "中文释义"

    if mode == "en_dictate_cn":
        v_key, b_key, is_dict = "word", "content", True
        visible_header_text = "可见内容"
        blank_header_text = "默写区域"
        if data_type == "irregular_verbs":
            visible_header_text = "原型"
            blank_header_text = "过去/过分/中文" # 用户要求
        elif data_type == "phrases":
            visible_header_text = "短语"
            blank_header_text = "默写区域"
    elif mode == "cn_dictate_en":
        v_key, b_key, is_dict = "content", "word", True
        visible_header_text = "可见内容"
        blank_header_text = "默写区域"
        if data_type == "phrases":
            visible_header_text = "翻译"
            blank_header_text = "默写区域"
        # 不规则动词表没有“看中默英”模式，所以这里不需要处理 irregular_verbs
    elif mode == "normal": # Explicitly handle "normal" mode
        # Defaults are already set for "words" type in "normal" mode
        if data_type == "irregular_verbs":
            visible_header_text = "原型"
            blank_header_text = "过去式/过去分词/中文" # 用户要求
        elif data_type == "phrases":
            visible_header_text = "短语"
            blank_header_text = "翻译"
        # If data_type is "words", it will use the initial default values.

    # 3. 【布局逻辑】计算列宽
    # 序号列固定为 1.2cm (维持现状)
    index_width_cm = 1.2
    page_total_width = 19.0  # A4 可用宽度

    if mode == "cn_dictate_en":
        # 中文默写英语模式：中文/发音列设窄，英语默写列设宽
        calc_word_width_cm = 3.5  # 可见列宽度 (中文+发音)
        explanation_width_cm = (page_total_width / 2) - index_width_cm - calc_word_width_cm # 默写列会变宽
    else:
        # 英语显示模式：采用动态包裹逻辑
        word_stats = []
        for item in data:
            raw_val = str(item.get(v_key, ""))
            segments = re.split(r'\(|（|或|\[', raw_val)
            current_max_seg_width = 0
            for seg in segments:
                length = sum(2 if '\u4e00' <= char <= '\u9fff' else 1 for char in seg)
                if length > current_max_seg_width:
                    current_max_seg_width = length
            word_stats.append({'max_seg_width': current_max_seg_width})

        max_segment_width = max([s['max_seg_width'] for s in word_stats]) if word_stats else 0
        calc_word_width_cm = max(2.5, min(6.0, max_segment_width * 0.165 + 0.6))
        explanation_width_cm = (page_total_width / 2) - index_width_cm - calc_word_width_cm

    col_widths = [Cm(index_width_cm), Cm(calc_word_width_cm), Cm(explanation_width_cm)] * 2

    # 4. 创建表格
    table = document.add_table(rows=1, cols=6)
    table.autofit = False
    table.style = 'Table Grid'

    # 设置表头
    headers = ["序号", visible_header_text, blank_header_text]
    hdr_cells = table.rows[0].cells
    for i in range(6):
        table.columns[i].width = col_widths[i]
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(headers[i % 3])
        set_run_font(run, FONT_SIZE_CONTENT, bold=True)

    # 5. 填充数据
    for i in range(0, len(data), 2):
        row_data = data[i:i+2]
        row = table.add_row()
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        row.height = Cm(0.7) # 稍微增加行高容纳两行文字

        for j in range(2):
            if j < len(row_data):
                item = row_data[j]
                idx_txt = str(i + j + 1)
                vis_raw = str(item.get(v_key, ""))
                blk_raw = str(item.get(b_key, ""))

                # 【核心逻辑修正】：针对不同模式应用不同的换行策略
                if mode == "cn_dictate_en":
                    # 中文默写英语模式：发音在第一行，中文在下一行
                    # 寻找音标结尾的 ] 符号并在其后换行
                    if "]" in vis_raw:
                        vis_raw = vis_raw.replace("]", "]\n")
                    else:
                        # 如果没有音标，尝试在第一个汉字前切分
                        vis_raw = re.sub(r'([^\x00-\xff])', r'\n\1', vis_raw, count=1)
                else:
                    # 英语显示模式：维持原来的 ( [ 或 换行
                    if any(x in vis_raw for x in ["(", "（", "或", "["]):
                        vis_raw = vis_raw.replace("(", "\n(").replace("（", "\n（").replace("或", "\n或").replace("[", "\n[")

                vis_txt = _normalize_full_width_to_half_width(vis_raw)
                blk_txt = "" if is_dict else _normalize_full_width_to_half_width(blk_raw)

                # 填充单元格
                for k, txt in enumerate([idx_txt, vis_txt, blk_txt]):
                    cell = row.cells[j*3 + k]
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.space_after = Pt(2)
                    p.paragraph_format.line_spacing = 1.0

                    # 序号居中，文字靠左
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if k == 0 else WD_ALIGN_PARAGRAPH.LEFT
                    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

                    # 物理换行逻辑
                    lines = txt.split('\n')
                    for line_idx, line in enumerate(lines):
                        if not line.strip(): continue
                        if line_idx > 0:
                            p.add_run().add_break()
                        run = p.add_run(line.strip())
                        set_run_font(run, FONT_SIZE_INDEX if k == 0 else FONT_SIZE_CONTENT)

    # 6. 细化边框
    for row in table.rows:
        for cell in row.cells:
            tcPr = cell._element.get_or_add_tcPr()
            tcBorders = OxmlElement('w:tcBorders')
            for b in ['top', 'left', 'bottom', 'right']:
                elm = OxmlElement(f'w:{b}')
                elm.set(qn('w:val'), 'single')
                elm.set(qn('w:sz'), '4')
                tcBorders.append(elm)
            tcPr.append(tcBorders)

    document.save(output_file)
    print(f"\nWord document generated (mode: {mode}, blank column width: {explanation_width_cm:.2f}cm): {output_file}\n")
